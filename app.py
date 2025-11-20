"""CLI applicatie om tabellen uit PDF bestanden te halen."""
from __future__ import annotations

import json
from pathlib import Path
from typing import List, Optional

from openpyxl import Workbook

import typer

from pdf_reader.engines.base import summarize_tables
from pdf_reader.table_extractor import TableExtractor

app = typer.Typer(help="Experimenteer met verschillende tabel engines voor PDF's")


@app.command()
def list_engines() -> None:
    """Toon beschikbare engines."""
    extractor = TableExtractor()
    typer.echo("Beschikbare engines:")
    for engine in extractor.engines:
        typer.echo(f"- {engine}")


@app.command()
def extract(
    pdf: Path = typer.Argument(..., exists=True, readable=True, help="Pad naar het PDF bestand"),
    engine: Optional[List[str]] = typer.Option(None, "--engine", "-e", help="Naam van de engine(s)"),
    output: Optional[Path] = typer.Option(None, "--output", "-o", help="Schrijf JSON resultaat naar dit bestand"),
    excel_output: Optional[Path] = typer.Option(
        None,
        "--excel",
        help="Schrijf een Excel bestand met de tabellen per engine",
    ),
    view_json: bool = typer.Option(
        False,
        "--view-json",
        help="Toon de volledige JSON output in de terminal",
    ),
    tune_pdfplumber: bool = typer.Option(
        False,
        "--tune-pdfplumber",
        help=(
            "Gebruik pymupdf4llm als referentie en pas pdfplumber-instellingen"
            " recursief aan om dichter bij het referentieresultaat te komen"
        ),
    ),
    tuning_depth: int = typer.Option(
        4,
        "--tuning-depth",
        min=0,
        help="Aantal recursiestappen voor het afstemmen van pdfplumber",
    ),
    min_words: Optional[List[int]] = typer.Option(
        None,
        "--min-words",
        help="Optionele lijst met min_words kandidaten voor het afstemmen",
    ),
) -> None:
    """Voer tabel extractie uit."""
    extractor = TableExtractor()
    results = extractor.extract(
        pdf,
        engines=engine,
        tune_pdfplumber=tune_pdfplumber,
        tuning_depth=tuning_depth,
        min_word_options=min_words,
    )
    # We'll also print summary to console
    for result in results:
        typer.secho(f"Resultaten voor {result.engine}", fg=typer.colors.GREEN, bold=True)
        if not result.tables:
            typer.echo("Geen tabellen gevonden\n")
            continue
        typer.echo(summarize_tables(result.tables))
        typer.echo("")
    data = {
        result.engine: [
            [[cell.text for cell in row] for row in table]
            for table in result.tables
        ]
        for result in results
    }

    if output:
        output.write_text(json.dumps(data, ensure_ascii=False, indent=2))
        typer.echo(f"Resultaat opgeslagen in {output}")

    if excel_output:
        workbook = Workbook()
        # Remove the default empty sheet
        default_sheet = workbook.active
        workbook.remove(default_sheet)
        for engine_name, tables in data.items():
            worksheet = workbook.create_sheet(engine_name[:31] or "result")
            if not tables:
                worksheet.append(["Geen tabellen gevonden"])
                continue
            for idx, table in enumerate(tables, start=1):
                worksheet.append([f"Table {idx}"])
                for row in table:
                    worksheet.append(row)
                worksheet.append([])  # lege regel tussen tabellen
        workbook.save(excel_output)
        typer.echo(f"Excel bestand opgeslagen in {excel_output}")

    if view_json:
        typer.echo("Volledige JSON output:")
        typer.echo(json.dumps(data, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    app()
