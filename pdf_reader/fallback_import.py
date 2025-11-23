"""Fallback flow voor tabelimport tussen pymupdf4llm en pdfplumber."""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterable, Sequence, Tuple

from openpyxl import Workbook

from .engines.base import Table
from .table_extractor import TableExtractor

DEFAULT_ENGINE_ORDER: Tuple[str, ...] = ("pymupdf4llm", "pdfplumber")


@dataclass
class FallbackImportResult:
    engine: str
    tables: list[Table]
    markdown: str


def _pad_row(row: Sequence[str], width: int) -> Sequence[str]:
    if len(row) >= width:
        return row
    return list(row) + [""] * (width - len(row))


def table_to_markdown(table: Table) -> str:
    """Render een tabel als markdown."""

    if not table:
        raise ValueError("Lege tabel kan niet naar markdown worden omgezet")

    text_rows = [[cell.text for cell in row] for row in table]
    column_count = max(len(row) for row in text_rows)

    def format_row(row: Sequence[str]) -> str:
        padded = _pad_row(row, column_count)
        return "| " + " | ".join(padded) + " |"

    header_divider = "| " + " | ".join(["---"] * column_count) + " |"
    lines = [format_row(text_rows[0]), header_divider]
    for row in text_rows[1:]:
        lines.append(format_row(row))
    return "\n".join(lines)


def import_table_with_fallback(
    pdf_path: Path | str,
    engine_order: Iterable[str] | None = None,
    min_rows: int | None = None,
    extractor: TableExtractor | None = None,
) -> FallbackImportResult:
    """
    Probeer een PDF te importeren met configurabele enginevolgorde.

    De standaard volgorde is ``pymupdf4llm`` gevolgd door ``pdfplumber``. Per
    engine wordt gekeken of er minimaal één tabel is. Als ``min_rows`` is
    opgegeven, moet de eerste tabel minimaal zoveel rijen bevatten; zo niet, dan
    wordt de volgende engine geprobeerd. Als geen enkele engine voldoet, wordt
    ``ImportError`` geworpen.

    Returns
    -------
    FallbackImportResult
        Bevat de gekozen engine, de gevonden tabel en de markdown-weergave
        daarvan.
    """

    order = list(engine_order) if engine_order is not None else list(DEFAULT_ENGINE_ORDER)
    extractor = extractor or TableExtractor()
    pdf_path = Path(pdf_path)

    for engine_name in order:
        try:
            results = extractor.extract(pdf_path, engines=[engine_name])
        except ValueError as exc:
            raise ImportError(f"Onbekende engine opgegeven: {engine_name}") from exc

        if not results:
            continue

        result = results[0]
        if result.error:
            continue

        tables = result.tables
        if not tables:
            continue

        first_table = tables[0]
        if min_rows is not None and len(first_table) < min_rows:
            continue

        markdown = table_to_markdown(first_table)
        return FallbackImportResult(result.engine, tables, markdown)

    raise ImportError(
        "Geen geschikte engine gevonden die een tabel met voldoende rijen kon extraheren"
    )


def import_directory_with_fallback(
    input_dir: Path | str,
    output_excel: Path | str,
    engine_order: Iterable[str] | None = None,
    min_rows: int | None = None,
    extractor: TableExtractor | None = None,
) -> Dict[str, FallbackImportResult | ImportError]:
    """
    Verwerk alle PDF-bestanden in een map met de fallback helper en schrijf de
    resultaten naar één Excel-bestand.
    """

    input_path = Path(input_dir)
    output_path = Path(output_excel)
    extractor = extractor or TableExtractor()

    pdf_files = sorted(file for file in input_path.iterdir() if file.suffix.lower() == ".pdf")
    if not pdf_files:
        raise ValueError(f"Geen PDF-bestanden gevonden in {input_path}")

    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    results: Dict[str, FallbackImportResult | ImportError] = {}

    def unique_sheet_name(base: str, existing: set[str]) -> str:
        name = base[:31] or "result"
        candidate = name
        suffix = 1
        while candidate in existing:
            suffix_str = f"_{suffix}"
            candidate = f"{name[: 31 - len(suffix_str)]}{suffix_str}"
            suffix += 1
        return candidate

    sheet_names: set[str] = set()

    for pdf in pdf_files:
        sheet_name = unique_sheet_name(pdf.stem, sheet_names)
        sheet_names.add(sheet_name)
        worksheet = workbook.create_sheet(sheet_name)

        try:
            result = import_table_with_fallback(
                pdf, engine_order=engine_order, min_rows=min_rows, extractor=extractor
            )
            results[pdf.name] = result
            worksheet.append(["Engine", result.engine])
            worksheet.append([])
            for idx, table in enumerate(result.tables, start=1):
                worksheet.append([f"Table {idx}"])
                for row in table:
                    worksheet.append([cell.text for cell in row])
                worksheet.append([])
        except ImportError as exc:
            results[pdf.name] = exc
            worksheet.append(["Fout", str(exc)])

    workbook.save(output_path)
    return results
